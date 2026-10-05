import streamlit as st
import psutil
import requests
import ollama
import os
import io
import sys
import contextlib
import time
import pandas as pd
from PIL import Image
import PyPDF2
from pptx import Presentation
from duckduckgo_search import DDGS

# ---------------------------------------------------------
# Page Setup & CSS
# ---------------------------------------------------------
st.set_page_config(page_title="Nexus AI — Refined Interface", page_icon="✨", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
    .stApp { background: radial-gradient(850px 550px at 72% -10%,rgba(143,124,255,.17),transparent 60%), radial-gradient(650px 480px at 10% 105%,rgba(121,230,255,.10),transparent 60%), #080b16 !important; color: #eef3ff !important; font-family: Inter, ui-sans-serif, system-ui, sans-serif !important; }
    header[data-testid="stHeader"] { display: none !important; }
    section[data-testid="stSidebar"] { background-color: rgba(16,21,34,.94) !important; border-right: 1px solid #273148 !important; min-width: 268px !important; }
    .topbar { display: flex; align-items: center; justify-content: space-between; padding: 15px 27px; border-bottom: 1px solid #273148; background: rgba(10,14,24,.72); backdrop-filter: blur(12px); margin: -90px -40px 20px -40px; z-index: 999; }
    .topbar-title { display: flex; align-items: center; gap: 10px; }
    .topbar-title h1 { font-size: 18px; margin: 0; font-weight: 750; color: #eef3ff; line-height: 1; }
    .topbar-state { font-size: 12px; color: #8994ad; padding-left: 2px; }
    .online-dot { width: 7px; height: 7px; border-radius: 50%; }
    .deploy-btn-group { display: flex; gap: 9px; }
    .top-btn { height: 36px; padding: 0 16px; border-radius: 9px; border: 1px solid #273148; background: #101522; color: #eef3ff; font-size: 13px; font-weight: 650; cursor: pointer; }
    .top-btn.deploy { border: 0; background: linear-gradient(120deg,#8f7cff,#79e6ff); color: #07101b; }
    .brand { display: flex; align-items: center; gap: 10px; font-size: 20px; font-weight: 750; letter-spacing: .2px; margin-bottom: 22px; color:#eef3ff;}
    .logo { width: 31px; height: 31px; border-radius: 9px; background: linear-gradient(135deg,#8f7cff,#79e6ff); display: grid; place-items: center; color: #07101b; box-shadow: 0 0 25px rgba(121,230,255,.12); }
    .logo svg { width: 20px; height: 20px; }
    div[data-testid="stSidebar"] .stButton > button { width: 100% !important; height: 43px !important; border: 1px solid #273148 !important; border-radius: 11px !important; background: #151b2b !important; color: #eef3ff !important; justify-content: center !important; padding: 0 13px !important; font-weight: 600 !important; font-size: 18px !important; transition: background 0.2s, border-color 0.2s; }
    div[data-testid="stSidebar"] .stButton > button:hover { border-color: #52627f !important; background: #182035 !important; }
    div[data-testid="stSidebar"] button[kind="primary"] { background: linear-gradient(120deg,#8f7cff,#79e6ff) !important; border-color: transparent !important; color: #07101b !important; box-shadow: 0 8px 25px rgba(121,230,255,.10) !important; justify-content: flex-start !important; font-size: 13.5px !important;}
    .status { margin-top: 18px; border: 1px solid #273148; border-radius: 14px; background: rgba(21,27,43,.72); padding: 14px 13px; }
    .status-title { display: flex; align-items: center; gap: 8px; color: #8994ad; font-size: 12px; font-weight: 700; text-transform: uppercase; letter-spacing: .08em; margin-bottom: 14px; }
    .metric { margin: 12px 0; }
    .metric-head { display: flex; justify-content: space-between; font-size: 12.5px; margin-bottom: 6px; color: #eef3ff; }
    .metric-head span { color: #8994ad; }
    .track { height: 5px; background: #252e43; border-radius: 20px; overflow: hidden; }
    .history-title { font-size: 12px; color: #8994ad; font-weight: 700; text-transform: uppercase; letter-spacing: .08em; margin: 0 6px 8px; }
    [data-testid="stChatMessage"] { padding: 13px 16px !important; border: 1px solid #273148 !important; border-radius: 15px !important; background: rgba(16,21,34,.92) !important; font-size: 14.5px !important; line-height: 1.58 !important; color: #e8edf8 !important; box-shadow: 0 8px 24px rgba(0,0,0,.08) !important; max-width: 800px; margin-bottom: 20px; }
    [data-testid="stChatMessage"]:nth-child(odd) { background: linear-gradient(135deg, rgba(143,124,255,.19), rgba(121,230,255,.08)) !important; border-color: rgba(143,124,255,.65) !important; }
    .stChatInputContainer { max-width: 850px !important; margin: 0 auto !important; border: 1px solid #2c3852 !important; background: rgba(16,21,34,.96) !important; border-radius: 18px !important; padding: 11px 12px 10px !important; box-shadow: 0 16px 45px rgba(0,0,0,.25) !important; }
    [data-testid="stSelectbox"] div[data-baseweb="select"] { height: 37px !important; min-height: 37px !important; border: 1px solid #273148 !important; border-radius: 10px !important; background: #151b2b !important; color: #eaf1ff !important; font-size: 12.5px !important; }
    [data-testid="stFileUploaderDropzone"] { border: 1px dashed #273148 !important; border-radius: 10px !important; background: #151b2b !important; padding: 10px !important; }
    .stTextArea textarea { background: #151b2b !important; color: #eaf1ff !important; border: 1px solid #273148 !important; border-radius: 10px !important; font-size: 12.5px !important; }
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# Telemetry Helpers & Core Tools
# ---------------------------------------------------------
def get_loaded_models():
    try:
        res = requests.get("http://localhost:11434/api/ps", timeout=1.5)
        if res.status_code == 200: return res.json().get("models", [])
    except: pass
    return []

def force_unload_all():
    for m in get_loaded_models():
        try: requests.post("http://localhost:11434/api/generate", json={"model": m.get("name"), "keep_alive": 0}, timeout=2.0)
        except: pass

def extract_text_from_multiple_files(uploaded_files):
    combined_text = ""
    for uploaded_file in uploaded_files:
        file_type = uploaded_file.name.split('.')[-1].lower()
        if file_type == 'txt':
            combined_text += f"\n--- {uploaded_file.name} ---\n" + uploaded_file.getvalue().decode("utf-8") + "\n"
        elif file_type == 'pdf':
            pdf_text = "".join([page.extract_text() or "" for page in PyPDF2.PdfReader(uploaded_file).pages])
            combined_text += f"\n--- {uploaded_file.name} ---\n" + pdf_text + "\n"
        elif file_type == 'csv':
            df = pd.read_csv(uploaded_file)
            combined_text += f"\n--- {uploaded_file.name} (Data Snapshot) ---\nColumns: {', '.join(df.columns)}\nPreview:\n{df.head(5).to_markdown()}\n"
    return combined_text

def create_presentation(title_text, points):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title_text
    slide.placeholders[1].text = "Generated autonomously by Nexus AI Core"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Key Insights & Overview"
    tf = slide2.placeholders[1].text_frame
    tf.text = points[0] if points else "Core concept generated by local model."
    for pt in points[1:]: tf.add_paragraph().text, tf.paragraphs[-1].level = pt, 0
    file_stream = io.BytesIO(); prs.save(file_stream); file_stream.seek(0)
    return file_stream

def fetch_web_search(query):
    try:
        results = DDGS().text(query, max_results=4)
        return "\n".join([f"- **{r['title']}**: {r['body']} ({r['href']})" for r in results])
    except Exception as e:
        return f"Web search failed: {str(e)}"

# LIVE DYNAMIC STATUS BADGE
try:
    requests.get("http://localhost:11434/", timeout=0.5)
    sys_status, sys_color = "online", "#57e39b"
except:
    sys_status, sys_color = "offline", "#ff6b6b"

# ---------------------------------------------------------
# Session State & Personas
# ---------------------------------------------------------
available_models = ["qwen2.5-coder:7b-instruct-q4_K_M", "llava:7b", "llama3.1", "deepseek-r1:1.5b"]
personas = {
    "💻 Code Wizard": "You are an expert software engineer and programming tutor. Provide clean, efficient code.",
    "🧮 Logic & Math Professor": "You are an advanced mathematical logic professor. Break down problems step-by-step.",
    "📈 Data Analyst": "You are an expert data scientist. When given data, provide precise insights or write Python pandas/matplotlib code to analyze it.",
    "✨ General Assistant": "You are Nexus AI, a helpful, precise, and balanced personal assistant."
}

if "active_model" not in st.session_state: st.session_state.active_model = available_models[0]
if "last_file_names" not in st.session_state: st.session_state.last_file_names = []
if "messages" not in st.session_state: st.session_state.messages = [{"role": "assistant", "content": "Nexus Core initialized. All systems ready."}]
if "active_tool" not in st.session_state: st.session_state.active_tool = "Chat"

# ---------------------------------------------------------
# Sidebar Layout 
# ---------------------------------------------------------
with st.sidebar:
    st.markdown("""<div class="brand"><div class="logo"><svg viewBox="0 0 24 24" fill="none"><path d="M12 2.5 20 7v10l-8 4.5L4 17V7l8-4.5Z" stroke="currentColor" stroke-width="1.6"/><circle cx="12" cy="12" r="2.1" fill="currentColor"/></svg></div><span>Nexus AI</span></div>""", unsafe_allow_html=True)
    
    if st.button("＋ New chat", use_container_width=True, type="primary"): 
        st.session_state.messages = [{"role": "assistant", "content": "Nexus Core initialized. All systems ready."}]
        st.session_state.active_tool = "Chat"
        st.rerun()
    
    if st.button("⌫ Purge memory", use_container_width=True): force_unload_all(); st.rerun()

    # MINI-LOGO TOOL SELECTOR
    st.markdown("<div class='history-title' style='margin-top:15px;'>Quick Tools</div>", unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    if c1.button("💬", help="Standard Chat"): st.session_state.active_tool = "Chat"; st.rerun()
    if c2.button("🌐", help="Web Search Mode"): st.session_state.active_tool = "Web Search"; st.rerun()
    if c3.button("📊", help="PowerPoint Mode"): st.session_state.active_tool = "PowerPoint"; st.rerun()
    
    mode_color = "#57e39b" if st.session_state.active_tool == "Chat" else "#79e6ff"
    st.markdown(f"<p style='font-size:12.5px; color:{mode_color}; margin-top:0px; margin-bottom: 10px; font-weight:600;'>Mode: {st.session_state.active_tool}</p>", unsafe_allow_html=True)

    st.markdown("<div class='history-title' style='margin-top:10px;'>Active Model Engine</div>", unsafe_allow_html=True)
    current_index = available_models.index(st.session_state.active_model) if st.session_state.active_model in available_models else 0
    new_selected_model = st.selectbox("Model Selector", available_models, index=current_index, label_visibility="collapsed")

    st.markdown("<div class='history-title' style='margin-top:15px;'>Document / Data Upload</div>", unsafe_allow_html=True)
    uploaded_files = st.file_uploader("Upload Files", type=['png', 'jpg', 'jpeg', 'pdf', 'txt', 'csv'], accept_multiple_files=True, label_visibility="collapsed")

    st.markdown("<div class='history-title' style='margin-top:15px;'>Active Persona</div>", unsafe_allow_html=True)
    selected_persona_name = st.selectbox("Persona Selector", list(personas.keys()), label_visibility="collapsed")
    current_system_prompt = personas[selected_persona_name]

    st.markdown("<div class='history-title' style='margin-top:10px;'>Additional Instructions</div>", unsafe_allow_html=True)
    custom_instructions = st.text_area("Custom Instructions", placeholder="e.g. 'Act like a pirate'...", height=68, label_visibility="collapsed")

    chat_export_markdown = "".join([f"**{'User' if m['role'] == 'user' else 'Nexus AI'}:** {m['content']}\n\n---\n" for m in st.session_state.messages])
    st.markdown("<div style='margin-top: 15px;'></div>", unsafe_allow_html=True)
    st.download_button(label="📥 Export Chat History", data=chat_export_markdown, file_name="nexus_chat.md", mime="text/markdown", use_container_width=True)

    mem, cpu = psutil.virtual_memory(), psutil.cpu_percent(interval=0.1)
    st.markdown(f"""
    <div class="status">
      <div class="status-title"><span class="dot"></span>System status</div>
      <div class="metric"><div class="metric-head"><span>RAM used</span><b>{(mem.total - mem.available) / (1024**3):.1f} / {mem.total / (1024**3):.0f} GB</b></div><div class="track"><i style="width:{mem.percent}%;background:#79e6ff;display:block;height:100%;border-radius:20px;"></i></div></div>
      <div class="metric"><div class="metric-head"><span>CPU usage</span><b>{cpu}%</b></div><div class="track"><i style="width:{cpu}%;background:#8f7cff;display:block;height:100%;border-radius:20px;"></i></div></div>
    </div>
    """, unsafe_allow_html=True)

# ---------------------------------------------------------
# Main Workspace & Topbar Header
# ---------------------------------------------------------
st.markdown(f"""
<header class="topbar">
  <div class="topbar-title"><span class="online-dot" style="background:{sys_color}; box-shadow: 0 0 9px {sys_color}88;"></span><h1>Nexus AI</h1><span class="topbar-state">local model · {sys_status}</span></div>
  <div class="deploy-btn-group"><button class="top-btn">Stop</button><button class="top-btn deploy">Deploy</button></div>
</header>
""", unsafe_allow_html=True)

for idx, msg in enumerate(st.session_state.messages):
    with st.chat_message(msg["role"], avatar="👤" if msg["role"] == "user" else "✨"):
        st.markdown(msg["content"])
        if "meta" in msg: st.markdown(f"<p style='font-size: 11.5px; color: #8994ad; margin-top: 5px; font-weight: 500;'>⏱️ {msg['meta']}</p>", unsafe_allow_html=True)
        if "ppt_data" in msg: st.download_button("📥 Download Presentation", data=msg["ppt_data"], file_name="Nexus_Presentation.pptx", key=f"ppt_{idx}")
        
        if msg["role"] == "assistant" and "```python" in msg["content"]:
            if st.button(f"▶ Run Code Snippet", key=f"run_{idx}"):
                try:
                    code_block = msg["content"].split("```python")[1].split("```")[0].strip()
                    st.markdown("**🛡️️ Execution Output:**")
                    buf = io.StringIO()
                    with contextlib.redirect_stdout(buf): exec(code_block, {})
                    res = buf.getvalue()
                    st.code(res, language="text") if res.strip() else st.success("Executed successfully with no output.")
                except Exception as ex: st.error(f"Error: {str(ex)}")

# ---------------------------------------------------------
# SMART ROUTER
# ---------------------------------------------------------
current_files = [f.name for f in uploaded_files] if uploaded_files else []
if current_files != st.session_state.last_file_names:
    st.session_state.last_file_names = current_files
    if uploaded_files:
        exts = [f.split('.')[-1].lower() for f in current_files]
        if any(e in ['png', 'jpg', 'jpeg'] for e in exts):
            st.session_state.active_model = "llava:7b"
        elif any(e in ['pdf', 'txt'] for e in exts):
            st.session_state.active_model = "llama3.1"
        elif any(e == 'csv' for e in exts):
            st.session_state.active_model = "qwen2.5-coder:7b-instruct-q4_K_M" 
    else:
        st.session_state.active_model = "qwen2.5-coder:7b-instruct-q4_K_M"
    st.session_state.messages = [{"role": "assistant", "content": "Nexus Core initialized. All systems ready."}]
    st.rerun()

if new_selected_model != st.session_state.active_model:
    st.session_state.active_model = new_selected_model
    st.session_state.messages = [{"role": "assistant", "content": "Nexus Core initialized. All systems ready."}]
    st.rerun()

# ---------------------------------------------------------
# Dynamic Chat Input
# ---------------------------------------------------------
input_placeholder = "Message Nexus AI..."
if st.session_state.active_tool == "Web Search": input_placeholder = "Search the web for..."
elif st.session_state.active_tool == "PowerPoint": input_placeholder = "Topic for presentation..."

prompt = st.chat_input(input_placeholder)

# ---------------------------------------------------------
# Generation Loop
# ---------------------------------------------------------
if prompt:
    # Determine the execution mode (auto-apply prefix if tool is active)
    is_ppt = prompt.lower().startswith("ppt:") or st.session_state.active_tool == "PowerPoint"
    is_search = prompt.lower().startswith("search:") or st.session_state.active_tool == "Web Search"

    user_msg = {"role": "user", "content": prompt}
    if uploaded_files:
        exts = [f.name.split('.')[-1].lower() for f in uploaded_files]
        if any(e in ['pdf', 'txt', 'csv'] for e in exts):
            user_msg["content"] = f"Documents/Data:\n{extract_text_from_multiple_files(uploaded_files)}\n\nQuery: {prompt}"
        elif any(e in ['png', 'jpg', 'jpeg'] for e in exts):
            img = Image.open(io.BytesIO([f for f in uploaded_files if f.name.split('.')[-1].lower() in ['png', 'jpg', 'jpeg']][0].getvalue()))
            img.thumbnail((1600, 1600))
            buf = io.BytesIO(); img.save(buf, format='JPEG')
            user_msg["images"] = [buf.getvalue()]

    st.session_state.messages.append(user_msg)
    with st.chat_message("user", avatar="👤"): st.markdown(prompt)

    with st.chat_message("assistant", avatar="✨"):
        response_box = st.empty()
        start_time = time.time()
        
        # INTERCEPT 1: POWERPOINT ENGINE
        if is_ppt:
            topic = prompt[4:].strip() if prompt.lower().startswith("ppt:") else prompt
            response_box.markdown(f"⚙️ **Generating presentation slides for:** *{topic}*...")
            res = ollama.chat(model="llama3.1", messages=[{"role": "user", "content": f"Provide 4 concise bullet points explaining: {topic}"}])
            pts = [line.strip("-* ").strip() for line in res["message"]["content"].split("\n") if line.strip()][:4]
            ppt_stream = create_presentation(topic, pts or ["Introduction", "Features", "Impact", "Outlook"])
            
            dur = round(time.time() - start_time, 1)
            full_text = f"✅ PowerPoint presentation generated for **{topic}**!"
            st.session_state.messages.append({"role": "assistant", "content": full_text, "ppt_data": ppt_stream.getvalue(), "meta": f"Rendered in {dur}s"})
            st.session_state.active_tool = "Chat" # Auto-reset to standard chat
            st.rerun()

        # INTERCEPT 2: WEB SEARCH ENGINE
        elif is_search:
            topic = prompt[7:].strip() if prompt.lower().startswith("search:") else prompt
            response_box.markdown(f"🌐 **Searching the live web for:** *{topic}*...")
            
            web_results = fetch_web_search(topic)
            search_sys_prompt = f"{current_system_prompt}\n\nUse the following real-time web search results to accurately answer the user's query about '{topic}'. Cite your sources if possible.\n\nWEB RESULTS:\n{web_results}"
            
            full_text = f"**Search Results Analyzed for:** *{topic}*\n\n"
            try:
                stream = ollama.chat(model=st.session_state.active_model, messages=[{"role": "system", "content": search_sys_prompt}, {"role": "user", "content": f"Based on the web results, explain: {topic}"}], stream=True)
                for chunk in stream:
                    full_text += chunk["message"]["content"]
                    response_box.markdown(full_text + " ▍")
                response_box.markdown(full_text)
                
                dur = round(time.time() - start_time, 1)
                st.session_state.messages.append({"role": "assistant", "content": full_text, "meta": f"{len(full_text.split())} words generated in {dur}s"})
                st.session_state.active_tool = "Chat" # Auto-reset to standard chat
                st.rerun()
            except Exception as e:
                st.error("Error generating search response."); st.code(str(e))

        # STANDARD LLM CHAT
        else:
            full_text = ""
            opts = {"temperature": 0.1, "repeat_penalty": 1.5} if "llava" in st.session_state.active_model else {"temperature": 0.6 if "deepseek" in st.session_state.active_model else 0.3}
            
            active_sys_prompt = current_system_prompt
            if custom_instructions.strip(): active_sys_prompt += f"\n\nUser Custom Instructions: {custom_instructions}"
            
            try:
                stream = ollama.chat(model=st.session_state.active_model, messages=[{"role": "system", "content": active_sys_prompt}] + st.session_state.messages, stream=True, keep_alive="1m", options=opts)
                for chunk in stream:
                    text_chunk = chunk["message"]["content"].replace("<think>", "🧠 **Thinking Process:**\n```text\n").replace("</think>", "\n```\n\n---\n**Final Answer:**\n\n")
                    full_text += text_chunk
                    response_box.markdown(full_text + " ▍")
                
                response_box.markdown(full_text)
                dur = round(time.time() - start_time, 1)
                st.session_state.messages.append({"role": "assistant", "content": full_text, "meta": f"{len(full_text.split())} words generated in {dur}s"})
                st.rerun()
            except Exception as e:
                st.error("Error connecting to Ollama."); st.code(str(e))